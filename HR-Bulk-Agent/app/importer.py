import csv
from io import BytesIO, StringIO
from typing import Any

from fastapi import UploadFile

from app.models import BulkImportRow, ChangeEvent, EmployeePayload, Platform
from app.validation import validate_change_request


HEADER_MAP = {
    "사번": "employee_id",
    "이름": "name",
    "재직상태": "employment_status",
    "입사일": "hire_date",
    "퇴사일": "termination_date",
    "휴직시작일": "leave_start_date",
    "복직일": "leave_end_date",
    "부서": "department",
    "직급": "rank",
    "직책": "title",
    "고용형태": "employment_type",
    "회사": "company",
    "법인": "company",
    "이메일": "email",
    "휴대전화": "phone",
    "전화번호": "phone",
    "생년월일": "birth_date",
    "변경유형": "event_type",
    "대상플랫폼": "target_platforms",
}


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for key, value in row.items():
        mapped = HEADER_MAP.get(str(key).strip(), str(key).strip())
        normalized[mapped] = str(value).strip() if value is not None else None
    return normalized


def parse_platforms(value: str | None) -> list[Platform]:
    if not value:
        return [Platform.gift24, Platform.workd]
    platforms: list[Platform] = []
    for item in value.replace(";", ",").split(","):
        name = item.strip()
        if name:
            platforms.append(Platform(name))
    return platforms or [Platform.gift24, Platform.workd]


async def parse_upload(file: UploadFile) -> list[dict[str, Any]]:
    content = await file.read()
    filename = file.filename or ""
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(cell is not None for cell in row)]
    text = content.decode("utf-8-sig")
    return list(csv.DictReader(StringIO(text)))


async def preview_import(file: UploadFile) -> list[dict[str, Any]]:
    rows = await parse_upload(file)
    previews: list[dict[str, Any]] = []
    for index, raw in enumerate(rows, start=2):
        normalized = normalize_row(raw)
        event_type = ChangeEvent(normalized.pop("event_type", "hire"))
        target_platforms = parse_platforms(normalized.pop("target_platforms", None))
        employee = EmployeePayload(**normalized)
        change = BulkImportRow(
            row_number=index,
            event_type=event_type,
            employee=employee,
            target_platforms=target_platforms,
        )
        previews.append(
            {
                "row_number": index,
                "change": change.model_dump(mode="json"),
                "validation_errors": validate_change_request(change),
            }
        )
    return previews

